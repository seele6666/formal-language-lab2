"""Convert the experiment report markdown into a Word document."""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Inches

HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
IMAGE_RE = re.compile(r"^!\[(.*)\]\((.*)\)$")


def add_code_block(document: Document, lines: list[str]) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run("\n".join(lines))
    run.font.name = "Consolas"


def convert(markdown_path: Path, output_path: Path) -> None:
    document = Document()
    base_dir = markdown_path.parent
    lines = markdown_path.read_text(encoding="utf-8").splitlines()
    index = 0
    in_code = False
    code_lines: list[str] = []
    code_lang = ""

    while index < len(lines):
        line = lines[index]

        if line.startswith("```"):
            if not in_code:
                in_code = True
                code_lang = line[3:].strip()
                code_lines = []
            else:
                in_code = False
                add_code_block(document, code_lines)
                if code_lang == "text":
                    pass
                code_lines = []
                code_lang = ""
            index += 1
            continue

        if in_code:
            code_lines.append(line)
            index += 1
            continue

        heading_match = HEADING_RE.match(line)
        if heading_match:
            level = len(heading_match.group(1))
            document.add_heading(heading_match.group(2), level=min(level, 4))
            index += 1
            continue

        image_match = IMAGE_RE.match(line)
        if image_match:
            image_path = base_dir / image_match.group(2)
            if image_path.exists():
                document.add_paragraph(image_match.group(1))
                document.add_picture(str(image_path), width=Inches(6.2))
            else:
                document.add_paragraph(f"[缺少图片] {image_path}")
            index += 1
            continue

        if line.startswith("- "):
            document.add_paragraph(line[2:], style="List Bullet")
            index += 1
            continue

        if not line.strip():
            index += 1
            continue

        document.add_paragraph(line)
        index += 1

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))


def main() -> None:
    if len(sys.argv) != 3:
        raise SystemExit("Usage: py scripts/md_to_docx.py <input.md> <output.docx>")
    convert(Path(sys.argv[1]), Path(sys.argv[2]))


if __name__ == "__main__":
    main()
